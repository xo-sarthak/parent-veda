# =============================================================================
#  content_overview_xlsx.py -- the commissioning overview, as a spreadsheet
# -----------------------------------------------------------------------------
#      python tools/content_overview_xlsx.py
#
#  The PDFs list every slot in the app, down to each rail card's artwork. That
#  is the right level for building the thing and the wrong level for deciding
#  WHO TO HIRE. This is the other view: one row per commissionable piece --
#  a film, an article, a course, a carousel -- with the expert each one needs.
#
#  ⚠️ MICRO-COPY IS EXCLUDED ON PURPOSE. Card titles, tab labels, hero lines and
#  section headings are written by whoever builds the screen, not commissioned
#  from a specialist, and 154 rows of TTC tile copy would bury the 52 articles
#  that actually need an andrologist.
#
#  ⚠️ SO ARE FAQs AND RED-FLAG CALLOUTS, because they are PART OF the read they
#  sit in. A doctor writing "Whose side is it, really" writes its four
#  questions and its when-to-see-someone box in the same sitting; listing them
#  separately would treble the apparent number of commissions.
#
#  ⚠️ THE EXPERT COLUMN IS A PROPOSAL, NOT A FACT READ OUT OF THE CODE. It is
#  derived from the area a piece sits in, and the rules are all in EXPERTS
#  below so they can be argued with in one place.
# =============================================================================

import csv
import io
import os
import re
from collections import Counter, OrderedDict, defaultdict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(ROOT, 'research', 'content-inventory')
OUT = os.path.join(ROOT, 'docs', 'content-brief', 'CONTENT-OVERVIEW.xlsx')

STAGES = OrderedDict([
    ('ttc', 'Trying to conceive'),
    ('preg', 'Pregnancy'),
    ('pp', 'Parenting'),
])

# Micro-copy and navigation text: written by the build, not commissioned.
SKIP_TYPES = {
    'tile copy (title + blurb)', 'hub copy', 'bracket copy', 'closing line',
    'faq', 'red-flag callout', 'tracker-copy', 'section-hero copy',
    'area title + blurb', 'hero-copy', 'journey-config', 'section-page',
    'tool question', 'checker copy (questions + outcomes)',
}

# A readable name for each raw content_type.
TYPE_NAME = {
    'read/article': 'Article', 'article': 'Article', 'short-article': 'Article',
    'read': 'Article', 'read (article)': 'Article',
    'guide': 'Guide', 'day-by-day guide': 'Guide', 'ceremony guide': 'Guide',
    'route/decision page': 'Guide', 'scan guide': 'Guide', 'scan extras': 'Guide',
    'video-script': 'Video', 'video script + "why watch" blurb': 'Video',
    'film slot copy': 'Video', 'podcast episode outline': 'Podcast',
    'myth-fact': 'Myth vs fact', 'carousel-card-copy': 'Carousel',
    'cards page': 'Carousel', 'infographic-copy': 'Infographic',
    'chart card page': 'Chart', 'chart': 'Chart', 'comparison table': 'Comparison table',
    'offering-blurb': 'Course or consult', 'class description': 'Course or class',
    'class description + 2 reviews': 'Course or class',
    'course/masterclass description': 'Course or masterclass',
    'practice-script': 'Practice script (to record)',
    'SOS step': 'Practice script (to record)',
    'affirmation': 'Practice script (to record)',
    'audio track copy': 'Audio track', 'ritual copy': 'Ritual',
    'garbh prompt': 'Garbh Sanskar prompt', 'garbh story': 'Story',
    'passage to read aloud': 'Story', 'bedtime/folk story': 'Story',
    'tale': 'Story', 'recipe': 'Recipe', 'recipe (full)': 'Recipe',
    'recipe (mother)': 'Recipe', 'food entry': 'Food entry',
    'diet chart': 'Diet chart', 'diet-chart copy': 'Diet chart',
    'can-i entry': 'Is-it-safe entry', 'condition explainer': 'Condition explainer',
    'test explainer': 'Test explainer', 'finding explainer': 'Report finding',
    'report finding': 'Report finding', 'symptom entry': 'Symptom entry',
    'symptom group': 'Symptom entry', 'skin page': 'Skin page',
    'ingredient entry': 'Ingredient entry', 'craving entry': 'Craving entry',
    'vaccine explainer': 'Vaccine explainer',
    'vaccine learn-why + after-care': 'Vaccine explainer',
    'milestone copy': 'Milestone', 'milestone': 'Milestone',
    'activity': 'Activity', 'step-list': 'Step list',
    'journey-day': 'Journey day', 'development area journey': 'Journey',
    'concern (guided questions + result)': 'Concern explainer',
    'remedy': 'Home remedy', 'name-entry': 'Baby name',
    'name-story': 'Name story', 'numerology-profile': 'Numerology profile',
    'script (what to say)': 'What-to-say script', 'flagged callout': 'Callout',
    'tradition content': 'Tradition reading', 'showcase Q and A': 'Ask Veda example',
    'product-blurb': 'Product write-up',
    'product "why" + pros/cons': 'Product write-up',
    'recommendation ("why" + consider + best-for + benefits)': 'Recommendation',
    'precheck question': 'Checklist item', 'mission copy': 'Partner mission',
    'trimester tip': 'Tips set', 'body-change copy': 'Body-change copy',
    'week-development copy': 'Weekly copy', 'week-content (9 fields)': 'Weekly content',
    'week-content, Devanagari': 'Weekly content (Hindi)',
    'week-content, full variant': 'Weekly content',
    'daily-moment (6 fields per day)': 'Daily content',
    'daily-moment fallback': 'Daily content',
    'daily-insight': 'Daily insight', 'journal prompt': 'Journal prompt',
    'nutrition note': 'Nutrition note', 'readiness copy': 'Guide',
    'tool-outcome copy': 'Tool result copy', 'screener question': 'Screener',
    'preview copy': 'Preview copy',
}

# -----------------------------------------------------------------------------
#  Who writes it. Matched on the AREA a piece sits in, most specific first.
# -----------------------------------------------------------------------------
EXPERTS = [
    # (stage or None, area substring, expert, second reviewer)
    ('ttc', 'His side', 'Andrologist', ''),
    ('ttc', 'PCOS', 'Gynaecologist with a PCOS or endocrine focus', ''),
    ('ttc', 'IVF', 'Fertility specialist (IVF)', ''),
    ('ttc', 'After a loss', 'Obstetrician', 'Bereavement counsellor'),
    ('ttc', 'Mind and body', 'Perinatal psychologist or counsellor', ''),
    ('ttc', 'Practices', 'Yoga or meditation teacher', 'Perinatal psychologist'),
    ('ttc', 'Test library', 'Pathologist or fertility specialist', ''),
    ('ttc', 'Vaccines', 'Physician (adult immunisation)', ''),
    ('ttc', 'Shop', 'Pharmacist or dietitian (evidence check)', ''),
    ('ttc', 'Paid services', 'Depends on the offering — see title', ''),
    ('ttc', 'Partner', 'Andrologist', 'Counsellor'),
    ('ttc', 'Can I', 'Gynaecologist', 'Pharmacist'),
    ('ttc', 'Films', 'Depends on the film — see title', ''),
    ('ttc', 'Daily home', 'Fertility specialist', 'Copywriter'),
    ('ttc', None, 'Gynaecologist (fertility)', ''),

    ('preg', 'Mind and mood', 'Perinatal psychologist', ''),
    ('preg', 'Belly and skin', 'Dermatologist', 'Obstetrician'),
    ('preg', 'Nutrition', 'Prenatal dietitian', ''),
    ('preg', 'Cravings', 'Prenatal dietitian', ''),
    ('preg', 'Spiritual reading', 'A scholar of that tradition specifically', ''),
    ('preg', 'Garbh Sanskar', 'Garbh Sanskar practitioner', 'Obstetrician'),
    ('preg', 'Read to baby', "Children's writer", ''),
    ('preg', 'Conditions', 'Obstetrician', ''),
    ('preg', 'Tests and scans', 'Obstetrician or radiologist', ''),
    ('preg', 'Report findings', 'Obstetrician', ''),
    ('preg', 'Symptoms', 'Obstetrician', ''),
    ('preg', 'Can I', 'Obstetrician', 'Pharmacist'),
    ('preg', 'Ready for birth', 'Midwife or obstetric nurse', ''),
    ('preg', 'Paid services', 'Depends on the offering — see title', ''),
    ('preg', 'Products', 'Product reviewer', 'Obstetrician for any claim'),
    ('preg', 'Journey map', 'Obstetrician', ''),
    ('preg', 'Weekly', 'Obstetrician', ''),
    ('preg', 'Daily home', 'Obstetrician', 'Copywriter'),
    ('preg', 'Ask Veda', 'Obstetrician', ''),
    ('preg', None, 'Obstetrician', ''),

    ('pp', 'Feeding', 'Lactation consultant', 'Paediatric dietitian'),
    ('pp', 'Health', 'Paediatrician', ''),
    ('pp', 'Vaccination', 'Paediatrician', ''),
    ('pp', 'Sleep', 'Paediatric sleep specialist', 'Paediatrician'),
    ('pp', 'Behaviour', 'Child psychologist', ''),
    ('pp', 'Potty', 'Paediatrician', 'Child psychologist'),
    ('pp', 'You, Maa', 'Postpartum specialist', 'Perinatal psychologist'),
    ('pp', 'You Maa', 'Postpartum specialist', 'Perinatal psychologist'),
    ('pp', 'First 40 Days', 'Postpartum doula or obstetrician', ''),
    ('pp', 'Jaapa', 'Postpartum doula or obstetrician', ''),
    ('pp', 'Early Learning', 'Child development specialist', ''),
    ('pp', 'Development', 'Child development specialist', ''),
    ('pp', 'Milestone', 'Paediatrician', 'Child development specialist'),
    ('pp', 'Leap', 'Child development specialist', ''),
    ('pp', 'Traditions', 'Cultural expert — regional, not one national voice', ''),
    ('pp', 'Nuskhe', 'Ayurvedic practitioner', 'Paediatrician (safety veto)'),
    ('pp', 'Food', 'Paediatric dietitian', ''),
    ('pp', 'Recipes', 'Paediatric dietitian', 'Recipe developer'),
    ('pp', 'Nutrition', 'Paediatric dietitian', ''),
    ('pp', 'names', 'Linguist or cultural expert', ''),
    ('pp', 'Naming', 'Linguist or cultural expert', ''),
    ('pp', 'Yoga', 'Postnatal yoga instructor', ''),
    ('pp', 'Watch', 'Depends on the film — see title', ''),
    ('pp', 'Recommendations', 'Product reviewer', 'Paediatrician for any claim'),
    ('pp', 'Products', 'Product reviewer', 'Paediatrician for any claim'),
    ('pp', 'Deals', 'Commercial team, not a content expert', ''),
    ('pp', 'Community', 'Moderator, not a writer', ''),
    ('pp', 'Find help', 'Directory research, not authored content', ''),
    ('pp', 'Courses', 'Depends on the course — see title', ''),
    ('pp', 'Journeys', 'Lactation consultant', 'Paediatrician'),
    ('pp', 'What Changed', 'Paediatrician', ''),
    ('pp', 'Grow', 'Child development specialist', ''),
    ('pp', 'Wallet', 'Administrative research (state-by-state)', ''),
    ('pp', 'Documents', 'Administrative research (state-by-state)', ''),
    ('pp', 'Sounds', 'Music producer', ''),
    ('pp', 'Scripts', 'Child psychologist', ''),
    ('pp', 'Ask Veda', 'Paediatrician', ''),
    ('pp', None, 'Paediatrician', ''),
]


# Areas that name a shelf rather than a subject. A film filed under "Films"
# tells you nothing about who should present it; its TAP PATH names the door,
# and the film catalogue often names the role outright.
GENERIC_AREAS = ('films', 'watch', 'hubs', 'courses (recorded lessons)',
                 'paid services', 'courses')

# ⚠️ THE TITLE IS USUALLY ENOUGH, AND WHERE IT IS, USE IT. "Male fertility
# consultation" needs an andrologist whatever shelf it sits on. This catches
# the offerings, courses and films whose area names a shelf rather than a
# subject, and it is the last thing tried so a real area rule always wins.
TITLE_HINTS = [
    (('semen', 'sperm', 'male fertility', 'androl', 'his side', 'testic'),
     'Andrologist'),
    (('gynaecolog', 'gynecolog'), 'Gynaecologist'),
    (('fertility specialist', 'fertility, honestly', 'couple assessment',
      'ninety-day', '90-day'), 'Fertility specialist'),
    (('pcos', 'insulin', 'hormonal'), 'Gynaecologist with a PCOS or endocrine focus'),
    (('ivf', 'iui', 'embryo', 'transfer', 'stimulat'), 'Fertility specialist (IVF)'),
    (('breastfeed', 'lactation', 'latch', 'milk supply', 'weaning'),
     'Lactation consultant'),
    (('sleep',), 'Paediatric sleep specialist'),
    (('tummy time', 'first words', 'coos', 'high-contrast', 'narrate',
      'play', 'milestone'), 'Child development specialist'),
    (('fever', 'rash', 'colic', 'teething', 'cough'), 'Paediatrician'),
    (('solids', 'iron-rich', 'first foods', 'puree'), 'Paediatric dietitian'),
    (('calm-down', 'tantrum', 'crying', 'soothe', 'white noise'),
     'Child psychologist'),
    (('for you, too', 'yourself', 'her fourth month', 'postpartum'),
     'Postpartum specialist'),
    (('pressure', 'expectation'), 'Psychologist or counsellor'),
    (('vaccin', 'immunis', 'immuniz'), 'Paediatrician'),
    (('yoga', 'asana', 'breathwork', 'pranayam'), 'Yoga instructor'),
    (('meditat', 'mindful', 'stillness'), 'Meditation teacher'),
    (('nutrition', 'diet', '食', 'recipe', '食物', 'eating', 'food'), 'Dietitian'),
    (('counsell', 'counsel', 'therapy', 'psycholog', 'grief', 'loss', 'anxiety',
      'depress', 'mental'), 'Psychologist or counsellor'),
    (('birth', 'labour', 'labor', 'delivery'), 'Childbirth educator'),
    (('miscarriage', 'bereave'), 'Obstetrician'),
    (('massage', 'physio'), 'Physiotherapist'),
    (('name', 'naming'), 'Linguist or cultural expert'),
    (('astrolog', 'numerolog'), 'Astrologer or numerologist'),
    (('tradition', 'ceremony', 'ritual', 'sanskar'), 'Cultural expert'),
]


def by_title(title):
    t = (title or '').lower()
    for keys, who in TITLE_HINTS:
        if any(k in t for k in keys):
            return who
    return None


def expert_for(stage, area, ctype, path='', notes='', title=''):
    a = area or ''
    # ⚠️ THE APP ALREADY SAYS WHO PRESENTS EACH FILM. `ttc_videos_data.dart`
    # carries an expert and an expertRole per slot, which the inventory copied
    # into the notes as "On camera: NAME, ROLE". Using it beats inferring from
    # the shelf a film happens to be filed under.
    m = re.search(r'On camera:\s*[^,]+,\s*([^.]+)', notes or '')
    if m:
        # "Gynaecologist · 14 years" -> "Gynaecologist". The years belong to a
        # person, and this column is a ROLE to go and hire.
        role = re.split(r'[·,]', m.group(1))[0].strip()
        role = re.sub(r'\s*\d+\s*years?$', '', role).strip()
        if role and 'role unstated' not in role.lower():
            return (role[0].upper() + role[1:], '')
    # A generic shelf: fall back to the door named in the tap path.
    if any(g in a.lower() for g in GENERIC_AREAS) and path:
        for st, key, who, second in EXPERTS:
            if st != stage or not key or who.startswith('Depends on'):
                continue          # a "Depends" rule here would beat the title
            if key.lower() in path.lower():
                return (who, second)
    if 'Story' in TYPE_NAME.get(ctype, '') or ctype in ('bedtime/folk story',):
        return ("Children's writer", '')
    for st, key, who, second in EXPERTS:
        if st != stage:
            continue
        if key is None or key.lower() in a.lower():
            if who.startswith('Depends on'):
                guess = by_title(title)
                if guess:
                    return (guess, '')
            return (who, second)
    return (by_title(title) or 'To be decided', '')


def num(s):
    s = (s or '').strip()
    return int(s) if s.isdigit() else 0


# ⚠️ "NEEDS A CLINICIAN" IS DERIVED FROM THE EXPERT, NOT FROM THE STATUS FIELD.
#
# The first cut of this sheet counted rows whose status was
# `clinical-review-owed` and reported **zero** for Parenting -- the stage
# carrying paediatric, feeding and safe-sleep content. The status vocabulary
# differs between stages because the inventories were produced by different
# passes, so counting on it compared two things that were never the same.
#
# Who has to sign a piece is a property of the piece, not of a label some
# earlier pass happened to attach. If the expert is a clinical role, a
# clinician signs it.
CLINICAL_WORDS = (
    'paediatric', 'pediatric', 'obstetric', 'gynaec', 'gynec', 'androlog',
    'dermatolog', 'psycholog', 'physician', 'dietitian', 'lactation',
    'patholog', 'radiolog', 'midwife', 'fertility specialist', 'doula',
    'sleep specialist', 'postpartum specialist', 'ayurvedic',
)


def is_clinical(expert, second):
    blob = ('%s %s' % (expert or '', second or '')).lower()
    return any(w in blob for w in CLINICAL_WORDS)


def rows_for(stage):
    out = []
    # ---- videos, from the image sheet (they carry the running time) --------
    seen = set()
    p = os.path.join(SRC, '%s_images.csv' % stage)
    for r in csv.DictReader(io.open(p, encoding='utf-8')):
        if r.get('kind') != 'video':
            continue
        # "Film for" before "Film", or the alternation eats the shorter one
        # first and leaves a stray "for" on the front of every title.
        # Only strip a real prefix. A bare "Film" was also matching the front
        # of "Film slot copy — ...", which is how a film to shoot ended up
        # titled "slot copy".
        title = re.sub(r'^(?:Film for|Video for|Poster frame for)\s*|'
                       r'^(?:Film|Video):\s*', '', r.get('item') or '').strip('" ')
        key = title.lower()
        if not title or key in seen:
            continue
        seen.add(key)
        who, second = expert_for(stage, r.get('area'), 'video-script',
                                 r.get('tap_path'), r.get('notes'), title)
        out.append([title, 'Video', r.get('tap_path'), r.get('area'), who, second,
                    'Yes' if is_clinical(who, second) else 'No',
                    'Not shot', 0, r.get('size_or_aspect') or '', r.get('source_file'),
                    r.get('notes') or ''])
    # ---- everything written ------------------------------------------------
    p = os.path.join(SRC, '%s_written.csv' % stage)
    for r in csv.DictReader(io.open(p, encoding='utf-8')):
        ct = r.get('content_type') or ''
        if ct in SKIP_TYPES:
            continue
        nice = TYPE_NAME.get(ct, ct[:1].upper() + ct[1:])
        if nice == 'Video':                      # already listed above
            continue
        who, second = expert_for(stage, r.get('area'), ct, r.get('tap_path'),
                                 '', r.get('title_or_id') or '')
        status = {'real': 'Written and shipping',
                  'clinical-review-owed': 'Draft — needs a clinician',
                  'coming-soon': 'Marked coming soon',
                  'derived-from-mother': 'Derived, no writing needed'}.get(
                      r.get('status'), 'Draft — placeholder')
        n = num(r.get('count')) or 1
        title = r.get('title_or_id') or ''
        if n > 1:
            title = '%s  (%d pieces)' % (title, n)
        out.append([title, nice, r.get('tap_path'), r.get('area'), who, second,
                    'Yes' if is_clinical(who, second) else 'No',
                    status, num(r.get('current_words')), '', r.get('source_file'),
                    r.get('purpose') or ''])
    return out


# -----------------------------------------------------------------------------
#  Workbook
# -----------------------------------------------------------------------------
HEAD = ['Title', 'Type', 'Where it is in the app', 'Area / door',
        'Expert needed', 'Second reviewer', 'Clinical sign-off?', 'Status today',
        'Words now', 'Length', 'In code', 'What it is for']
WIDTHS = [46, 20, 52, 26, 34, 28, 15, 24, 9, 14, 34, 60]

INK = '1F1A29'
HDR_FILL = PatternFill('solid', fgColor='2D144C')
BAND = PatternFill('solid', fgColor='F4F1F7')
REAL = PatternFill('solid', fgColor='E7F3EC')
THIN = Side(style='thin', color='DDD8E3')
BORDER = Border(bottom=THIN)


def sheet(wb, name, rows):
    ws = wb.create_sheet(name)
    ws.append(HEAD)
    for c in range(1, len(HEAD) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(name='Arial', size=9, bold=True, color='FFFFFF')
        cell.fill = HDR_FILL
        cell.alignment = Alignment(vertical='center', wrap_text=True)
        ws.column_dimensions[get_column_letter(c)].width = WIDTHS[c - 1]
    ws.row_dimensions[1].height = 30
    for i, r in enumerate(rows):
        ws.append(r)
        rr = i + 2
        for c in range(1, len(HEAD) + 1):
            cell = ws.cell(row=rr, column=c)
            cell.font = Font(name='Arial', size=9, color=INK)
            cell.alignment = Alignment(vertical='top', wrap_text=(c in (1, 3, 5, 12)))
            cell.border = BORDER
            if r[7].startswith('Written'):
                cell.fill = REAL
            elif i % 2:
                cell.fill = BAND
        ws.cell(row=rr, column=9).alignment = Alignment(horizontal='right', vertical='top')
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = 'A1:%s%d' % (get_column_letter(len(HEAD)), len(rows) + 1)
    return ws


def summary(wb, data):
    ws = wb.create_sheet('Overview', 0)
    ws.column_dimensions['A'].width = 44
    for col in 'BCDE':
        ws.column_dimensions[col].width = 17

    def title(row, txt, size=13):
        c = ws.cell(row=row, column=1, value=txt)
        c.font = Font(name='Arial', size=size, bold=True, color=INK)

    def head(row, cells):
        for i, t in enumerate(cells):
            c = ws.cell(row=row, column=i + 1, value=t)
            c.font = Font(name='Arial', size=9, bold=True, color='FFFFFF')
            c.fill = HDR_FILL
            c.alignment = Alignment(vertical='center', wrap_text=True)

    title(1, 'ParentVeda — content overview', 16)
    ws.cell(row=2, column=1,
            value='One row per commissionable piece, and the expert each one needs. '
                  'Card titles, tab labels and section headings are excluded: they are '
                  'written by the build, not commissioned.').font = Font(
        name='Arial', size=9, color='5A5265')
    ws.merge_cells('A2:E2')
    ws.row_dimensions[2].height = 28
    ws.cell(row=3, column=1,
            value='FAQs and red-flag callouts are counted inside the article they sit in, '
                  'because the same expert writes them in the same sitting.').font = Font(
        name='Arial', size=9, color='5A5265')
    ws.merge_cells('A3:E3')

    r = 5
    title(r, 'How much, per stage')
    r += 1
    head(r, ['Stage', 'Pieces', 'Needs clinical sign-off', 'Already written', 'Words in place'])
    r += 1
    first = r
    for st, nm in STAGES.items():
        rows = data[st]
        ws.cell(row=r, column=1, value=nm)
        ws.cell(row=r, column=2, value=len(rows))
        ws.cell(row=r, column=3, value=sum(1 for x in rows if x[6] == 'Yes'))
        ws.cell(row=r, column=4, value=sum(1 for x in rows if x[7].startswith('Written')))
        ws.cell(row=r, column=5, value=sum(x[8] for x in rows))
        r += 1
    ws.cell(row=r, column=1, value='All three')
    for c in range(2, 6):
        L = get_column_letter(c)
        ws.cell(row=r, column=c, value='=SUM(%s%d:%s%d)' % (L, first, L, r - 1))
    for rr in range(first, r + 1):
        for c in range(1, 6):
            cell = ws.cell(row=rr, column=c)
            cell.font = Font(name='Arial', size=9, bold=(rr == r), color=INK)
            cell.border = BORDER
            if c >= 2:
                cell.number_format = '#,##0'

    r += 3
    title(r, 'By kind of piece')
    r += 1
    head(r, ['Type'] + [STAGES[s] for s in STAGES] + ['Total'])
    r += 1
    kinds = Counter()
    per = {s: Counter(x[1] for x in data[s]) for s in STAGES}
    for s in STAGES:
        kinds.update(per[s])
    first = r
    for k, _ in kinds.most_common():
        ws.cell(row=r, column=1, value=k)
        for i, s in enumerate(STAGES):
            ws.cell(row=r, column=2 + i, value=per[s].get(k, 0))
        L1, L3 = get_column_letter(2), get_column_letter(1 + len(STAGES))
        ws.cell(row=r, column=2 + len(STAGES),
                value='=SUM(%s%d:%s%d)' % (L1, r, L3, r))
        r += 1
    for rr in range(first, r):
        for c in range(1, 2 + len(STAGES) + 1):
            cell = ws.cell(row=rr, column=c)
            cell.font = Font(name='Arial', size=9, color=INK)
            cell.border = BORDER
            if c >= 2:
                cell.number_format = '#,##0'

    r += 2
    title(r, 'By expert — who we have to find')
    r += 1
    head(r, ['Expert', 'Pieces', 'Stages'])
    r += 1
    who = Counter()
    where = defaultdict(set)
    for s in STAGES:
        for x in data[s]:
            who[x[4]] += 1
            where[x[4]].add(STAGES[s])
    first = r
    for k, n in who.most_common():
        ws.cell(row=r, column=1, value=k)
        ws.cell(row=r, column=2, value=n)
        ws.cell(row=r, column=3, value=', '.join(sorted(where[k])))
        r += 1
    for rr in range(first, r):
        for c in (1, 2, 3):
            cell = ws.cell(row=rr, column=c)
            cell.font = Font(name='Arial', size=9, color=INK)
            cell.border = BORDER
            if c == 2:
                cell.number_format = '#,##0'
    ws.freeze_panes = 'A5'
    return ws


if __name__ == '__main__':
    data = {s: rows_for(s) for s in STAGES}
    wb = Workbook()
    wb.remove(wb.active)
    summary(wb, data)
    for s, nm in STAGES.items():
        sheet(wb, nm, data[s])
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    wb.save(OUT)
    for s, nm in STAGES.items():
        print('%-22s %4d pieces  %4d need a clinician  %9s words in place'
              % (nm, len(data[s]), sum(1 for x in data[s] if x[6] == 'Yes'),
                 '{:,}'.format(sum(x[8] for x in data[s]))))
    print('-> %s' % os.path.relpath(OUT, ROOT))
