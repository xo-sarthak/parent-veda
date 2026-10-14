"""ParentVeda content tracker: what the prepared-content documents asked for,
what is in the app, what is not and why, what is needed. Short on purpose.
Writes <out>.pdf and <out>.xlsx."""
import sys
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle)

OUT = sys.argv[1]
DATE = '13 October 2026'

DONE, PART, NOT, NA = 'Done', 'Partly', 'Not done', 'Not started'

# (folder, item, what the document asked for, status, where to see it / why not, needed to finish, from whom)
ROWS = [
    # ---------------- Pregnancy week-wise content ----------------
    ('Pregnancy week-wise', 'Week pages, weeks 5–18 and 20–23 (18 docs)',
     'A full page for each week: size, about your baby, baby science cards, you this week, symptoms, diet, share with partner.',
     DONE, 'Pregnancy home → tap "Details" under the baby picture. Also opens from the "This week" and "About the size of" cards, More → Week by week, and the new-week message. Not in the doors.',
     '—', '—'),
    ('Pregnancy week-wise', 'Week 19',
     'The folder has no week 19 document (the week 20 doc says it was skipped on request).',
     NOT, 'Week 19 still shows the older, shorter page.', 'The week 19 document.', 'Content team'),
    ('Pregnancy week-wise', 'Weeks 1–4 and 24–40',
     'Not in the folder yet.', NOT, 'Those weeks show the older, shorter page.', 'The documents for these weeks.', 'Content team'),
    ('Pregnancy week-wise', 'Size comparison fruit',
     'Each doc names a fruit (raspberry, avocado…).', DONE,
     'Week page, home size line and journal all use the doc\'s fruit for weeks 5–23. Other weeks keep the Indian comparisons.', '—', '—'),
    ('Pregnancy week-wise', 'Symptom cards',
     'Docs say: "Symptom cards require medical review before publishing."', PART,
     'Shown now with the doctor line under them (your decision).', 'A clinician to review the symptom cards for weeks 5–23.', 'Medical reviewer'),
    ('Pregnancy week-wise', 'Week 22 swelling card',
     'Says "call promptly" without saying who to call (a preeclampsia warning sign).', PART,
     'Shown as written, under the heading "When to contact your doctor".', 'Change the doc line to "call your doctor promptly".', 'Content team'),
    ('Pregnancy week-wise', 'Weeks 5–8 partner scan lists',
     'The week 20 doc notes these still need updating to the "next four weeks" standard.', NOT,
     'Shown as the docs have them today.', 'Updated scan lists in the week 5–8 docs.', 'Content team'),
    ('Pregnancy week-wise', 'Ask Veda knowing the week pages',
     'Implied: the assistant should be able to answer from this content.', NOT,
     'Ask Veda does not know this text yet. Its knowledge lives in a separate code repository.', 'Access to the Ask Veda repository, then a re-import of the content.', 'Engineering'),

    # ---------------- My Child phase-wise content ----------------
    ('My Child phase-wise', 'Phase pages PH001–PH015 (15 docs)',
     'Each phase: a hero, "Living through this phase" cards with tips, a five-area child snapshot, and when to contact the doctor.',
     DONE, 'Parenting home: the phase name at the top, "This phase, explained", and "Read the full description" (opens the full doc page). Also the phase map\'s page and the "A new stage" message.',
     '—', '—'),
    ('My Child phase-wise', 'Child snapshot and milestone pages',
     'Five areas (Brain, Physical, Language, Emotional, Nutrition), each with a full milestone page.', DONE,
     'Parenting home → "How {name} is doing" → tap any of the five cards.', '—', '—'),
    ('My Child phase-wise', 'Activities (76)',
     '4–7 activities per phase with steps, things to remember, safety and variations.', DONE,
     'Parenting home → "Today" (the child\'s own phase comes first) → tap an activity. Also from each milestone page.',
     '—', '—'),
    ('My Child phase-wise', 'Activity length',
     'The docs give no duration or difficulty.', PART,
     'The app says "At your pace" rather than invent a number.', 'Add a duration to each activity in the docs, if wanted.', 'Content team'),
    ('My Child phase-wise', 'PH015 "Naming All the Feelings"',
     'Named on the Emotional milestone page, but the activity itself is not written in the doc.', NOT,
     'The milestone page shows the name without a link.', 'The written activity.', 'Content team'),
    ('My Child phase-wise', 'PH010 names',
     'Two milestones are named differently in the snapshot and on their pages.', PART,
     'Both shown as the doc has them.', 'Pick one name for each.', 'Content team'),
    ('My Child phase-wise', 'Paediatric sign-off',
     'Docs say: "Needs paediatric panel sign-off".', PART,
     'Shown now with their doctor lines (your decision).', 'Paediatric panel review of doctor lines, milestones and activities.', 'Medical reviewer'),
    ('My Child phase-wise', 'Phases 16–20 (19 months to 5 years)',
     'No documents yet.', NOT, 'These phases keep the older text.', 'The documents for phases 16–20.', 'Content team'),

    # ---------------- Garbh Sanskar ----------------
    ('Garbh Sanskar', "Mother's affirmation, every day (280)",
     'One line for each day of pregnancy for the mother to say to her baby.', DONE,
     'Pregnancy home → Garbh Sanskar → "Talking to your baby". Also Garbh door → Talk & read → "Today\'s passage to read aloud".', '—', '—'),
    ('Garbh Sanskar', "Father's script, every day (280)",
     'A matching line for the father each day.', DONE,
     'Dad mode → Today → "Read to your baby" card. Also Dad mode → Read aloud → Affirmations (top of the list).', '—', '—'),
    ('Garbh Sanskar', 'Story, every day (280)',
     'A short fable for each day, ending on a line to the baby.', DONE,
     'Garbh door → Talk & read → "Today\'s story". The Stories shelf lists every story reached so far, today first.', '—', '—'),
    ('Garbh Sanskar', 'Mantras and prayers (113)',
     'Hindu, Sikh, Muslim and Christian texts with their meaning.', DONE,
     'Garbh door → Talk & read → Spiritual reading → choose a faith → "Mantras and prayers" (or Shabads, Duas, Scripture).', '—', '—'),
    ('Garbh Sanskar', 'Sacred stories (70) and Gita conversations (180)',
     'Retold Hindu stories with a moral.', DONE,
     'Garbh door → Talk & read → Spiritual reading → Hinduism.', '—', '—'),
    ('Garbh Sanskar', 'Ragas and music (45 tracks)',
     'A listening plan by week, with health claims ("lowers cortisol", "528 Hz").', NOT,
     'No recordings exist, and the health claims break the app\'s rule never to claim a practice improves the baby or the body.',
     'The 45 recordings. The brief is in docs/GARBH-AUDIO-OWED.md.', 'Audio production'),
    ('Garbh Sanskar', '"Developmental focus", "Intent", "Prenatal intent" columns',
     'Lines such as "deep cellular protection" or "for the baby\'s brain development".', NOT,
     'Left out on purpose: they are health claims (your decision).', '—', '—'),
    ('Garbh Sanskar', 'Mantras given without their text (14)',
     'Seven Hindu and six Muslim entries say "(Full hymn)" or "(Recitation of…)"; Philippians 4 has no text.', NOT,
     'Not shown until the text is given.', 'The actual texts.', 'Content team'),
    ('Garbh Sanskar', 'Bible verses',
     'Written in the NIV translation, which is copyrighted.', PART,
     'Shown in the World English Bible (free to use) instead.', 'Confirm the WEB wording, or provide an NIV licence.', 'Product / legal'),
    ('Garbh Sanskar', 'Tone of the daily lines',
     '"Perfect" or "perfectly" appears 335 times and "unshakeable" 77 times; some lines promise "you are safe from all harm".', PART,
     'Shown as written. For a mother whose baby has a diagnosed condition, a daily promise of perfection can hurt.',
     'A decision: keep, or soften these words.', 'Product / content'),
    ('Garbh Sanskar', 'Faith review',
     'Implied by quoting scripture.', NOT, 'Not yet reviewed by someone of each faith.', 'A read-through per faith before launch.', 'Content team'),

    # ---------------- Book summary ----------------
    ('Book summaries', '40 book companions',
     'A full companion per book: what it is about, the author, core philosophy, 5 key ideas, chapter by chapter, ParentVeda\'s take, quotes, and a closing note.',
     DONE, 'Pregnancy: Read Next → Books (pregnancy and newborn books). Parenting: Reading → Book Summaries (parenting and newborn books). Both open the Book Companion reader.',
     '—', '—'),
    ('Book summaries', 'ParentVeda\'s Take as pointers',
     'Docs mark each point with ❤️ (good) or ⚠️ (worth knowing).', DONE,
     'Shown with the app\'s own drawn marks (heart-in-hand, flag), not emoji.', '—', '—'),
    ('Book summaries', 'Closing note (doctor / extra help / Indian context)',
     'Most books end with one of these.', DONE, 'A new last section in the companion, with its own tab.', '—', '—'),
    ('Book summaries', '"Garbh Sanskar: Ancient Secrets to Give Birth to Genius"',
     'Listed as book 14 in Volume Three\'s contents.', NOT,
     'The companion was never written into the file; only its name appears.', 'The written companion.', 'Content team'),
    ('Book summaries', 'Volume Four, Part 1 (books 1–12)',
     'Part 2 says it continues from Part 1.', NOT, 'Part 1 is not in the folder.', 'The Part 1 file.', 'Content team'),
    ('Book summaries', 'Book covers',
     'Docs say "(add cover art asset)".', NOT, 'Each book shows the typographic cover the reader already draws.', 'Cover images, with rights.', 'Design / content'),
    ('Book summaries', 'What to Expect When You\'re Expecting',
     'The new doc differs from the app\'s older companion, which also had Hindi.', PART,
     'The new doc shows (your decision); the old companion is kept in the code for revert. Hindi readers see English for this book.',
     'A Hindi translation, if wanted.', 'Content team'),
    ('Book summaries', 'Several books have no quotes',
     'Docs say quotes could not be verified under the copyright limit.', PART,
     'The quotes section is simply left out for those books.', '—', '—'),

    # ---------------- Pregnancy FAQs ----------------
    ('Pregnancy FAQs', '326 FAQs as Ask Veda\'s knowledge (Months 1–8)',
     'An FAQ library for the Ask Veda assistant, with answer, why, what it means, doctor line and takeaways.',
     PART, 'Built in the app and Ask Veda, and loaded into the database (13 Oct). Ask Veda reads each FAQ as fact and '
     'writes the answer for the mother\'s own week; it now runs on a larger free model that does this reliably. '
     'Answers go live when the Ask Veda server runs again (Railway is paused).',
     'Restart the Ask Veda server.', 'Engineering'),
    ('Pregnancy FAQs', 'Each FAQ as a page',
     'Not asked by the docs; your decision.', DONE,
     'Ask Veda → an answer\'s "More information" card opens the FAQ as a full page, with its related questions below.',
     '—', '—'),
    ('Pregnancy FAQs', 'Suggested questions by month',
     'Not asked by the docs; your decision.', DONE,
     'Ask Veda home: the Pregnancy card shows questions from her month; "You could also ask" too.', '—', '—'),
    ('Pregnancy FAQs', '39 FAQs held for clinical sign-off',
     'Docs mark them "Clinical review: Yes", most "holds for Ishaan\'s sign-off".', NOT,
     'Held out of answers, pages and suggestions (your decision). Bleeding, miscarriage, preeclampsia, results, '
     'reduced movements, delivery mode, preterm.', 'Sign-off, one FAQ at a time.', 'Ishaan / clinical reviewer'),
    ('Pregnancy FAQs', 'Month 9 (weeks 36–40)', 'The library stops at Month 8.', NOT,
     'Weeks 36+ see Month 8\'s questions.', 'The Month 9 doc.', 'Content team'),
    ('Pregnancy FAQs', 'Related questions', 'Each FAQ lists 3 related questions by title.', PART,
     'Matched to real FAQs; 8 titles name a question that does not exist, so they are left off.',
     'Fix the 8 titles (listed in STILL-OPEN §88).', 'Content team'),
    ('Pregnancy FAQs', 'Verification notes', 'Each FAQ ends with sources and checks.', PART,
     'Not shown to mothers: they are writers\' notes (search volumes, "No discrepancies"). They stay in the docs.',
     '—', '—'),

    # ---------------- Not started ----------------
    # ---------------- Products ----------------
    ('Products', '373 pregnancy, hospital-bag and nursery picks',
     'Catalogue: 45 shelves of Budget and Premium picks, what to look for, and a week to buy by.', PART,
     'Products → Pregnancy: 12 new categories, plus 6 shelves inside the store\'s own categories, where they lead. '
     'Shown as ParentVeda\'s shortlist: "Price varies", no buy button, a plain cover (your decision). '
     '"For you right now" offers each near its buy-by week.',
     'Prices, buy links and photos per product.', 'Product / content team'),
    ('Products', '281 toys, scored', 'Two toy sheets: score, age, skills, budget.', PART,
     'Products → Parenting → Play & Development: ranked by score, the score on every card and page '
     '(your decision), part-scores where the sheet has them.',
     'Prices, links, photos; ages for the 145 toys from the Products Master, which has none.', 'Product / content team'),
    ('Products', 'Toy sheet clean-up', 'Ids and verification in the toy sheets.', PART,
     '30 rows reuse ids (PV-CRA-001..015 for both foil-art kits and crayons); 83 toys are marked not yet verified; '
     'some rows look like product and brand are swapped. Shown as the sheet has them.',
     'Renumber the ids; check the flagged rows.', 'Product team'),
    ('Products', 'Fetal dopplers and herbal anti-nausea teas (20)', 'Two catalogue shelves.', NOT,
     'Held for clinical review (your decision): home dopplers can falsely reassure; the app advises against '
     'herbal mixes without a doctor.', 'A clinical decision on both shelves.', 'Ishaan / clinical reviewer'),
    ('Products', 'Pregnancy books in the catalogue', 'Master Library books.', DONE,
     'A book that has a ParentVeda companion opens the companion instead of a product page.', '—', '—'),
    ('Products', '285 books to read with your child', 'Book Recommendations sheet: 291 rows.', DONE,
     'Reading (parenting) → "Books to read with your child": by the child\'s age, goal chips, Essentials first. '
     '14 rows with shifted columns were repaired; 6 duplicates merged.',
     'Fill the empty scoring sheets if scores should show.', 'Content team'),
    ('Products', 'New Product Categories sheet', 'A fifth sheet of toys.', NA,
     'Every row is already in the 0–5 master, so it adds nothing.', '—', '—'),
    ('Not started', 'Recipes, Instagram scraping', 'Both folders are empty.', NA, '—', 'The files.', 'Content team'),

    # ---------------- Checks ----------------
    ('Checks', 'Plagiarism check', 'Requested: check the documents are original.', NOT,
     'Not run (skipped for now). About 200 web searches, roughly US$2 (about ₹170).', 'A go-ahead to run it.', 'Product'),
    ('Checks', 'Warm, comforting tone', 'Requested: the app\'s own words should be warm and human.', DONE,
     'Four of our own labels made warmer. Document text left as written; cold or worrying lines listed above.', '—', '—'),
    ('Checks', "The app's latest marks", "Requested: use the app's current marks and glyphs.", DONE,
     'Drawn marks replace Material icons and emoji on the week page, phase pages, activities, Read Next and books.', '—', '—'),
]

STATUS_COLOR = {DONE: '#2E7D4F', PART: '#B7791F', NOT: '#B4443C', NA: '#6B6B6B'}

# The decision log (the user, 2026-10-13: "mention the questions that I am
# answering, the questions that you are asking"). (date, job, question, answer)
DECISIONS = [
    ('13 Oct', 'Weeks', 'Where should the week content live?',
     'On the week Details page (pregnancy home → Details under the baby), in the docs\' own sections.'),
    ('13 Oct', 'Weeks', 'Docs compare size to Western fruit; the app used Indian items. Which?', 'Use the doc\'s fruit.'),
    ('13 Oct', 'Weeks', 'Docs say symptom cards need medical review. Ship them?', 'Ship now with the doctor footer; review owed.'),
    ('13 Oct', 'Weeks', 'Week 5: revised doc or the older file?', 'Revised doc wins; the older file is kept.'),
    ('13 Oct', 'Weeks', '(You asked) Where can these changes be seen? Are the doors affected?',
     'Details pill, the "This week" and size cards, More → Week by week, the new-week message. Not the doors.'),
    ('13 Oct', 'All', '(Your rule) Use Mobbin for screens and a good UX.', 'Followed for every new screen part.'),
    ('13 Oct', 'Next', 'What to take next from the folder?', 'My Child phase-wise content.'),
    ('13 Oct', 'Phases', '(Your rule) The document is the source of truth; phases are the parenting side.', 'Applied.'),
    ('13 Oct', 'Phases', 'Doc phase names or the app\'s?', 'The doc\'s names.'),
    ('13 Oct', 'Phases', '"How your child is doing": only what is changing, or the doc\'s 5 domains?', 'The doc\'s 5 domains.'),
    ('13 Oct', 'Phases', 'What happens to the older phase text?', 'Doc wins on screen; the older text is kept underneath.'),
    ('13 Oct', 'Phases', 'Where do the doc activities go?', 'Into Today, the child\'s own phase first.'),
    ('13 Oct', 'Drive', 'Can the Drive folder be read directly?', 'Yes, once shared; new files are pulled from Drive.'),
    ('13 Oct', 'Garbh', 'Where does the day-by-day sheet go?', 'Mother\'s line, father\'s line and the story, each by day.'),
    ('13 Oct', 'Garbh', 'Ragas have no audio and make health claims. What then?', 'Wait for the audio; brief logged.'),
    ('13 Oct', 'Garbh', 'Mantras quote scripture (and NIV is copyrighted). How?',
     'Text and meaning, no health claims; Bible verses in a free translation.'),
    ('13 Oct', 'Garbh', 'Hindu sacred stories: where?', 'A new section in the Hindu spiritual library.'),
    ('13 Oct', 'Review', '(Your request) Recheck all work, check plagiarism, warm tone, keep a tracker.', 'Done, except the plagiarism check.'),
    ('13 Oct', 'Review', 'How long should the tracker be?', 'Short: 5–8 pages, not the 100+ page gap reports.'),
    ('13 Oct', 'Review', 'Run the plagiarism check (about US$2 / ₹170)?', 'Skip for now.'),
    ('13 Oct', 'Books', 'Where should the 41 book companions go?', 'By stage: pregnancy shelf and parenting shelf.'),
    ('13 Oct', 'Books', 'Include the "Genius" Garbh Sanskar book?', 'Include, with its critique (but the companion is missing from the file).'),
    ('13 Oct', 'Books', 'New What to Expect doc vs the older bilingual companion?', 'New doc wins; the old one is kept for revert.'),
    ('13 Oct', 'All', '(Your rule) Marks and glyphs: use the app\'s latest ones.',
     'Applied: drawn marks replace Material icons and emoji on the week page, phase pages, activities, Read Next and books.'),
    ('13 Oct', 'All', '(Your rule) Keep the questions and answers in the tracker.', 'This page.'),
    ('13 Oct', 'Check', '(You asked) After the compaction, is everything from the Drive so far in and maintained?',
     'Yes: re-running all four converters changed nothing; Drive unchanged since 11 Oct; 122 tests passed.'),
    ('13 Oct', 'Next', '(You asked) What is next in the Drive?', 'FAQs, with access to the Ask Veda repository.'),
    ('13 Oct', 'FAQs', 'Should Ask Veda show the FAQ\'s own answer word for word on a close match, or always write fresh?',
     'Always write fresh for where she is (her week); the FAQs are facts, not templates. AI cost is fine.'),
    ('13 Oct', 'FAQs', 'Where else should the FAQs show?', 'Ask Veda\'s suggested questions by her month, and each FAQ as a page.'),
    ('13 Oct', 'FAQs', 'What about FAQs marked for clinical review?', 'Hold them out until signed off (39 FAQs).'),
    ('13 Oct', 'FAQs', 'Who loads the FAQs into the live database?', 'You run it; commands handed over.'),
    ('13 Oct', 'FAQs', '(You asked) What is the live load, and can git run first?',
     'Export, import and embedding into Supabase; commit any time, push the Ask Veda repo only after the load.'),
    ('13 Oct', 'FAQs', 'The small free model refused week-28 folic acid; Haiku has no credit. Switch Groq models?',
     'Yes: gpt-oss-120b, same free-tier limits (1,000 requests a day, 8,000 tokens a minute).'),
    ('13 Oct', 'Products', 'No price, link or photo in the sheets. How do products appear?',
     'As ParentVeda\'s shortlist: brand, product, Budget/Premium tag, what to look for, "Price varies", no buy button.'),
    ('13 Oct', 'Products', 'The store\'s made-up products where the sheets cover the same thing?',
     'Sheets lead; the older products stay below. Nothing removed.'),
    ('13 Oct', 'Products', 'Show the toy scores to parents?', 'Rank by them and show them: "essential to make judgment".'),
    ('13 Oct', 'Products', 'Where do the 291 children\'s books go?', 'A parenting Reading shelf, by the child\'s age and goal.'),
    ('13 Oct', 'Products', 'Fetal dopplers and herbal teas clash with the app\'s clinical advice. What then?',
     'Hold both for clinical review.'),
    ('13 Oct', 'All', '(Your rule) Keep the tracker document maintained for review.', 'Updated after every job.'),
]


def pdf():
    ss = getSampleStyleSheet()
    body = ParagraphStyle('b', parent=ss['Normal'], fontName='Helvetica', fontSize=8.2, leading=10.4)
    head = ParagraphStyle('h', parent=body, fontName='Helvetica-Bold', textColor=colors.white)
    h1 = ParagraphStyle('h1', parent=ss['Title'], fontName='Helvetica-Bold', fontSize=20, leading=24, alignment=TA_LEFT)
    h2 = ParagraphStyle('h2', parent=ss['Heading2'], fontName='Helvetica-Bold', fontSize=13, spaceBefore=4, spaceAfter=6)
    lead = ParagraphStyle('l', parent=body, fontSize=10, leading=14)

    doc = SimpleDocTemplate(OUT + '.pdf', pagesize=landscape(A4), leftMargin=14 * mm, rightMargin=14 * mm,
                            topMargin=12 * mm, bottomMargin=12 * mm, title='ParentVeda content tracker')
    st = []
    st.append(Paragraph('ParentVeda — prepared content tracker', h1))
    st.append(Paragraph(f'What the documents in the shared Drive folder asked for, what is in the app, what is not and why, '
                        f'and what is needed to finish. Updated {DATE}.', lead))
    st.append(Spacer(1, 8))

    folders = ['Pregnancy week-wise', 'My Child phase-wise', 'Garbh Sanskar', 'Book summaries', 'Pregnancy FAQs', 'Products', 'Not started', 'Checks']
    summary = [[Paragraph(x, head) for x in ('Folder', 'Done', 'Partly', 'Not done / not started', 'In one line')]]
    lines = {
        'Pregnancy week-wise': '18 of the docs are live on the week page. Week 19 and weeks 1–4, 24–40 are owed.',
        'My Child phase-wise': 'All 15 docs are live on the parenting home. Phases 16–20 are owed.',
        'Garbh Sanskar': 'Daily lines, stories and prayers are live. Ragas wait for audio.',
        'Book summaries': '40 of 41 companions are live on both shelves. One is missing from its file; Part 1 of Volume Four is owed.',
        'Pregnancy FAQs': '326 of 365 loaded into Ask Veda; live when its server runs again. 39 wait for clinical sign-off.',
        'Products': '654 products and 285 kids\' books in the app, as the team\'s shortlist. Prices, links and photos owed; 20 held.',
        'Not started': 'Recipes and Instagram folders are empty.',
        'Checks': 'Tone done. Plagiarism check not run yet.',
    }
    for f in folders:
        rs = [r for r in ROWS if r[0] == f]
        summary.append([Paragraph(f'<b>{f}</b>', body),
                        Paragraph(str(sum(r[3] == DONE for r in rs)), body),
                        Paragraph(str(sum(r[3] == PART for r in rs)), body),
                        Paragraph(str(sum(r[3] in (NOT, NA) for r in rs)), body),
                        Paragraph(lines[f], body)])
    t = Table(summary, colWidths=[45 * mm, 18 * mm, 18 * mm, 40 * mm, 145 * mm])
    t.setStyle(TableStyle([('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#2F2C30')),
                           ('GRID', (0, 0), (-1, -1), 0.4, colors.HexColor('#D9D4DC')),
                           ('VALIGN', (0, 0), (-1, -1), 'TOP'), ('TOPPADDING', (0, 0), (-1, -1), 5),
                           ('BOTTOMPADDING', (0, 0), (-1, -1), 5)]))
    st.append(t)
    st.append(Spacer(1, 10))
    st.append(Paragraph('How to read this', h2))
    for line in ['<b>Done</b>: in the app now, as the document asked.',
                 '<b>Partly</b>: in the app, with a difference you should know about (the row says what).',
                 '<b>Not done</b>: not in the app; the row says why and what would finish it.',
                 '"Where to see it" is the tap path in the app. Everything listed is English only, by the current policy.']:
        st.append(Paragraph('• ' + line, lead))

    widths = [52 * mm, 62 * mm, 19 * mm, 74 * mm, 38 * mm, 24 * mm]
    for f in folders:
        st.append(PageBreak())
        st.append(Paragraph(f, h2))
        data = [[Paragraph(x, head) for x in ('Item', 'What the document asked for', 'Status',
                                              'Where to see it / why not', 'Needed to finish', 'From')]]
        style = [('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#2F2C30')),
                 ('GRID', (0, 0), (-1, -1), 0.4, colors.HexColor('#D9D4DC')),
                 ('VALIGN', (0, 0), (-1, -1), 'TOP'), ('TOPPADDING', (0, 0), (-1, -1), 4),
                 ('BOTTOMPADDING', (0, 0), (-1, -1), 4)]
        for i, r in enumerate([r for r in ROWS if r[0] == f], start=1):
            data.append([Paragraph(f'<b>{r[1]}</b>', body), Paragraph(r[2], body),
                         Paragraph(f'<font color="{STATUS_COLOR[r[3]]}"><b>{r[3]}</b></font>', body),
                         Paragraph(r[4], body), Paragraph(r[5], body), Paragraph(r[6], body)])
            if i % 2 == 0:
                style.append(('BACKGROUND', (0, i), (-1, i), colors.HexColor('#F7F5F8')))
        t = Table(data, colWidths=widths, repeatRows=1)
        t.setStyle(TableStyle(style))
        st.append(t)

    st.append(PageBreak())
    st.append(Paragraph('What we are waiting for', h2))
    owed = [[Paragraph(x, head) for x in ('From', 'What')]]
    by = {}
    for r in ROWS:
        if r[5] not in ('—', ''):
            by.setdefault(r[6], []).append(f'{r[5]} <i>({r[1]})</i>')
    for who, items in by.items():
        owed.append([Paragraph(f'<b>{who}</b>', body), Paragraph('<br/>'.join('• ' + x for x in items), body)])
    t = Table(owed, colWidths=[45 * mm, 224 * mm])
    t.setStyle(TableStyle([('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#2F2C30')),
                           ('GRID', (0, 0), (-1, -1), 0.4, colors.HexColor('#D9D4DC')),
                           ('VALIGN', (0, 0), (-1, -1), 'TOP'), ('TOPPADDING', (0, 0), (-1, -1), 5),
                           ('BOTTOMPADDING', (0, 0), (-1, -1), 5)]))
    st.append(t)

    st.append(PageBreak())
    st.append(Paragraph('Decisions: the questions asked and your answers', h2))
    st.append(Paragraph('Every choice that shaped what is in the app, in the order it was made.', lead))
    st.append(Spacer(1, 6))
    dec = [[Paragraph(x, head) for x in ('Date', 'Job', 'Question', 'Answer')]]
    dstyle = [('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#2F2C30')),
              ('GRID', (0, 0), (-1, -1), 0.4, colors.HexColor('#D9D4DC')),
              ('VALIGN', (0, 0), (-1, -1), 'TOP'), ('TOPPADDING', (0, 0), (-1, -1), 4),
              ('BOTTOMPADDING', (0, 0), (-1, -1), 4)]
    for i, (date, job, qn, ans) in enumerate(DECISIONS, start=1):
        dec.append([Paragraph(date, body), Paragraph(f'<b>{job}</b>', body), Paragraph(qn, body), Paragraph(ans, body)])
        if i % 2 == 0:
            dstyle.append(('BACKGROUND', (0, i), (-1, i), colors.HexColor('#F7F5F8')))
    t = Table(dec, colWidths=[16 * mm, 20 * mm, 118 * mm, 115 * mm], repeatRows=1)
    t.setStyle(TableStyle(dstyle))
    st.append(t)

    def footer(c, d):
        c.saveState()
        c.setFont('Helvetica', 7.5)
        c.setFillColor(colors.HexColor('#8B8494'))
        c.drawString(14 * mm, 7 * mm, f'ParentVeda content tracker · {DATE}')
        c.drawRightString(landscape(A4)[0] - 14 * mm, 7 * mm, f'Page {d.page}')
        c.restoreState()

    doc.build(st, onFirstPage=footer, onLaterPages=footer)


def xlsx():
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = 'Tracker'
    cols = ['Folder', 'Item', 'What the document asked for', 'Status', 'Where to see it / why not',
            'Needed to finish', 'From']
    ws.append(cols)
    for r in ROWS:
        ws.append(list(r))
    widths = [20, 38, 55, 12, 70, 40, 20]
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = w
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor='2F2C30')
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical='top')
        row[3].font = Font(bold=True, color=STATUS_COLOR[row[3].value].lstrip('#'))
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions

    wd = wb.create_sheet('Decisions')
    wd.append(['Date', 'Job', 'Question', 'Answer'])
    for row in DECISIONS:
        wd.append(list(row))
    for i, w in enumerate([10, 12, 70, 70]):
        wd.column_dimensions[chr(65 + i)].width = w
    for c in wd[1]:
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor='2F2C30')
    for row in wd.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical='top')
    wd.freeze_panes = 'A2'
    wb.save(OUT + '.xlsx')


pdf()
xlsx()
print('ok', len(ROWS), 'rows')
