"""Write the output.

The workbook is the deliverable: `parentveda_questions.xlsx` with
  - Summary          counts per stage and topic
  - All questions    every question, one row each, with stage / topic /
                     cluster / canonical / source / language / URL
  - Ranked · <stage> the clustered, demand-ranked list for each stage
The same tables also land as CSV (ranked_<stage>.csv, ranked_all.csv,
question_pool.csv) so they diff in git and load anywhere, and a README.md
shows the top of each list without opening a spreadsheet.
"""

from __future__ import annotations

import csv
import time
from collections import Counter
from pathlib import Path

from .cluster import Cluster
from .schema import STAGES, UNCLEAR

STAGE_LABEL = {"ttc": "Trying to Conceive", "pregnancy": "Pregnancy",
               "parenting": "Parenting", "skilling": "Skilling", UNCLEAR: "Unclear"}

RANKED_COLS = ["rank", "stage", "topic", "cluster_id", "canonical_question", "cluster_size", "total_hits",
               "source_spread", "languages", "new_this_run", "top_stems", "sample_phrasings"]
POOL_COLS = ["stage", "topic", "cluster_id", "cluster_rank", "cluster_size", "canonical_question",
             "question", "normalized", "source", "lang", "hits", "url", "first_seen_run"]


def _ranked_row(rank: int, c: Cluster) -> dict:
    return {
        "rank": rank, "stage": c.stage, "topic": c.topic, "cluster_id": c.cluster_id,
        "canonical_question": c.canonical, "cluster_size": c.size, "total_hits": c.hits,
        "source_spread": "; ".join(f"{s}:{n}" for s, n in c.sources.most_common()),
        "languages": "; ".join(f"{l}:{n}" for l, n in c.langs.most_common()),
        "new_this_run": "yes" if c.new_this_run else "",
        "top_stems": "; ".join(s for s, _ in c.stems.most_common(3)),
        "sample_phrasings": " | ".join(c.samples),
    }


def _pool_rows(clusters: dict[str, list[Cluster]]) -> list[dict]:
    pool: list[dict] = []
    for stage in list(STAGES) + [UNCLEAR]:
        for rank, c in enumerate(clusters.get(stage, []), start=1):
            for m in c.members:
                topic = m.topics.most_common(1)[0][0] if m.topics else c.topic
                for raw in m.raws:
                    pool.append({
                        "stage": stage, "topic": topic, "cluster_id": c.cluster_id, "cluster_rank": rank,
                        "cluster_size": c.size, "canonical_question": c.canonical, "question": raw,
                        "normalized": m.norm,
                        "source": "; ".join(f"{s}:{n}" for s, n in m.sources.most_common()),
                        "lang": m.langs.most_common(1)[0][0], "hits": m.hits,
                        "url": m.urls[0] if m.urls else "", "first_seen_run": m.first_seen,
                    })
    return pool


def write_all(out_dir: Path, clusters: dict[str, list[Cluster]], stats: dict) -> list[Path]:
    out_dir.mkdir(parents=True, exist_ok=True)
    written: list[Path] = []

    ranked_by_stage = {s: [_ranked_row(i, c) for i, c in enumerate(clusters.get(s, []), start=1)]
                       for s in list(STAGES) + [UNCLEAR]}
    pool = _pool_rows(clusters)

    # CSVs
    combined: list[dict] = []
    for stage, rows in ranked_by_stage.items():
        p = out_dir / f"ranked_{stage}.csv"
        _csv(p, RANKED_COLS, rows); written.append(p)
        combined.extend(rows)
    p = out_dir / "ranked_all.csv"; _csv(p, RANKED_COLS, combined); written.append(p)
    p = out_dir / "question_pool.csv"; _csv(p, POOL_COLS, pool); written.append(p)

    # Workbook
    try:
        p = out_dir / "parentveda_questions.xlsx"
        _xlsx(p, ranked_by_stage, pool, stats)
        written.append(p)
    except Exception as e:  # the CSVs are the record; the workbook is the convenience
        print(f"xlsx skipped: {e}")

    p = out_dir / "README.md"
    p.write_text(_readme(clusters, stats), encoding="utf-8"); written.append(p)
    return written


def _csv(path: Path, cols: list[str], rows: list[dict]) -> None:
    with path.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=cols)
        w.writeheader()
        for r in rows:
            w.writerow(r)


def _xlsx(path: Path, ranked_by_stage: dict[str, list[dict]], pool: list[dict], stats: dict) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="2F4F4F")

    def sheet(ws, cols, rows, widths):
        ws.append(cols)
        for cell in ws[1]:
            cell.font = head_font; cell.fill = head_fill
            cell.alignment = Alignment(vertical="center")
        for r in rows:
            ws.append([r.get(c, "") for c in cols])
        for j, c in enumerate(cols, start=1):
            ws.column_dimensions[get_column_letter(j)].width = widths.get(c, 12)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

    wb = Workbook()

    # Summary
    ws = wb.active; ws.title = "Summary"
    ws.append(["ParentVeda question research pool"]); ws["A1"].font = Font(bold=True, size=14)
    ws.append([f"Generated {time.strftime('%Y-%m-%d %H:%M')} · run #{stats.get('run_id')}"])
    ws.append(["Questions only, never answers. Ranked by distinct phrasings per cluster."])
    ws.append([])
    ws.append(["Pipeline counts"]); ws.cell(ws.max_row, 1).font = Font(bold=True)
    for k, v in stats.get("counts", {}).items():
        ws.append([k, v])
    ws.append([])
    ws.append(["Stage", "Questions", "Clusters"]);
    for cell in ws[ws.max_row]: cell.font = Font(bold=True)
    stage_q = Counter(r["stage"] for r in pool)
    for s in list(STAGES) + [UNCLEAR]:
        ws.append([STAGE_LABEL[s], stage_q.get(s, 0), len(ranked_by_stage.get(s, []))])
    ws.append([])
    ws.append(["Stage", "Topic", "Questions", "Clusters"])
    for cell in ws[ws.max_row]: cell.font = Font(bold=True)
    tq = Counter((r["stage"], r["topic"]) for r in pool)
    tc = Counter((r["stage"], r["topic"]) for rows in ranked_by_stage.values() for r in rows)
    for s in list(STAGES) + [UNCLEAR]:
        for (st, t), n in sorted(((k, v) for k, v in tq.items() if k[0] == s), key=lambda kv: -kv[1]):
            ws.append([STAGE_LABEL[st], t, n, tc.get((st, t), 0)])
    ws.append([])
    ws.append(["Source", "Questions"])
    for cell in ws[ws.max_row]: cell.font = Font(bold=True)
    sq = Counter()
    for r in pool:
        for part in r["source"].split("; "):
            sq[part.split(":")[0]] += 1
    for src, n in sq.most_common():
        ws.append([src, n])
    ws.column_dimensions["A"].width = 44; ws.column_dimensions["B"].width = 26
    ws.column_dimensions["C"].width = 12; ws.column_dimensions["D"].width = 12

    # All questions
    sheet(wb.create_sheet("All questions"), POOL_COLS, pool,
          {"question": 80, "normalized": 60, "canonical_question": 60, "source": 22, "topic": 20,
           "url": 40, "stage": 11})

    # Ranked per stage
    for s in list(STAGES) + [UNCLEAR]:
        sheet(wb.create_sheet(f"Ranked · {STAGE_LABEL[s]}"[:31]), RANKED_COLS, ranked_by_stage.get(s, []),
              {"canonical_question": 70, "source_spread": 30, "sample_phrasings": 110, "top_stems": 30,
               "languages": 16, "topic": 20, "stage": 11})
    wb.save(path)


def _readme(clusters: dict[str, list[Cluster]], stats: dict, top_n: int = 30) -> str:
    L = []
    L.append("# Question research pool\n")
    L.append("What real people ask, by ParentVeda stage, ranked by how many distinct ways it was asked.\n")
    L.append("Questions only, never answers. Generated by `tools/question_miner/` — see its README "
             "for sources and method. Re-run `python tools/question_miner/run.py run` to refresh.\n")
    L.append(f"Generated {time.strftime('%Y-%m-%d %H:%M')}. Run #{stats.get('run_id')}.\n")
    L.append("## Numbers\n")
    L.append("| | count |\n|---|---|")
    for k, v in stats.get("counts", {}).items():
        L.append(f"| {k} | {v} |")
    L.append("")
    L.append("| stage | questions (distinct phrasings) | clusters | in top cluster |\n|---|---|---|---|")
    for stage in list(STAGES) + [UNCLEAR]:
        cs = clusters.get(stage, [])
        n = sum(c.size for c in cs)
        L.append(f"| {STAGE_LABEL[stage]} | {n} | {len(cs)} | {cs[0].size if cs else 0} |")
    L.append("")
    L.append("## Files\n")
    L.append("- `parentveda_questions.xlsx` — the workbook: Summary, All questions (every question with stage, "
             "topic, cluster, source, language, URL), and one ranked sheet per stage.")
    L.append("- `ranked_<stage>.csv` / `ranked_all.csv` — the ranked lists as CSV.")
    L.append("- `question_pool.csv` — every question with its tags, as CSV.")
    L.append("- `questions.sqlite` — the raw store (gitignored; rebuilt by a run).\n")
    for stage in list(STAGES) + [UNCLEAR]:
        cs = clusters.get(stage, [])
        if not cs:
            continue
        L.append(f"## {STAGE_LABEL[stage]} — top {min(top_n, len(cs))} of {len(cs)}\n")
        L.append("| # | question | topic | asked | sources | e.g. |\n|---|---|---|---|---|---|")
        for i, c in enumerate(cs[:top_n], start=1):
            src = ", ".join(f"{s} {n}" for s, n in c.sources.most_common(3))
            eg = " / ".join(s.replace("|", "-") for s in c.samples[:2])
            L.append(f"| {i} | {c.canonical.replace('|', '-')} | {c.topic} | {c.size} | {src} | {eg} |")
        L.append("")
    return "\n".join(L)
